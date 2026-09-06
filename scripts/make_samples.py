"""生成评测/演示用示例文档：1 份 PDF（员工手册）、1 份 Word（产品说明书）、2 份 TXT。

文档内容为虚构场景（星辰科技），覆盖真实文档的常见结构：
多级标题、正文长段落、表格、FAQ 问答体 —— 用来检验解析与切分策略。
"""

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

OUT = ROOT / "data" / "samples"
OUT.mkdir(parents=True, exist_ok=True)


***REMOVED*** ==================== 1. 员工手册 PDF ====================

HANDBOOK = [
    ("第一章 公司简介", [
        "星辰科技成立于2018年，是一家专注于企业级人工智能应用的科技公司。公司总部位于杭州，在北京、上海设有分公司，目前员工规模约500人。"]),
    ("第二章 入职与试用期", []),
    ("2.1 入职材料", [
        "新员工办理入职手续时，需携带以下材料：身份证原件及复印件、学历学位证书原件、上一家单位的离职证明、三个月内的体检报告、本人银行卡（用于工资发放）。"]),
    ("2.2 试用期", [
        "试用期一般为3个月；管理岗位及核心技术岗位经审批可延长至6个月。试用期考核合格者按期转正，考核不合格的，公司将依法解除劳动合同。"]),
    ("第三章 考勤制度", []),
    ("3.1 工作时间", [
        "公司实行每周五天工作制，工作时间为9:00-18:00，其中12:00-13:30为午餐及休息时间。员工可在8:30-9:30之间弹性到岗，到岗后开始计算工时。"]),
    ("3.2 加班与调休", [
        "加班需提前在OA系统提交申请并获得直属上级审批。工作日加班按1:1比例调休，节假日加班按国家规定支付加班费或1:3调休。调休有效期三个月，逾期作废。"]),
    ("第四章 请假制度", []),
    ("4.1 年假", [
        "员工入职满1年不满3年的，每年享有5天带薪年假；满3年不满5年的，每年10天；满5年及以上的，每年15天。年假当年有效，最多可顺延至次年3月底。"]),
    ("4.2 病假", [
        "病假需提供二级及以上医院开具的诊断证明。病假期间工资按本人日工资的60%发放。连续病假超过3天的，需同时提供病历复印件。"]),
    ("4.3 事假", [
        "事假为无薪假，全年累计不超过15天，且每月不超过3天。事假需提前1个工作日申请，紧急情况可事后24小时内补办手续。"]),
    ("4.4 婚假、产假与丧假", [
        "婚假为10天，需在领证后一年内休完。产假为158天（含国家规定产假及地方奖励假）。直系亲属（父母、配偶、子女）去世，丧假为3天。"]),
    ("第五章 薪酬与福利", []),
    ("5.1 发薪日", [
        "公司于每月10日通过银行代发上月工资，遇节假日提前至最近的工作日发放。"]),
    ("5.2 社保与公积金", [
        "公司为员工缴纳五险一金，缴费基数为员工本人上年度月平均工资。"]),
    ("5.3 补充福利", [
        "公司为全体员工购买补充商业医疗保险，涵盖门诊与住院报销。每年组织一次全面体检。春节、中秋等节日发放节日礼品或礼品卡。每周五下午提供下午茶。"]),
    ("5.4 年终奖", [
        "年终奖根据公司年度经营情况与员工个人绩效综合评定，于次年春节前发放。具体金额不作承诺，不写入劳动合同。"]),
    ("第六章 费用报销", []),
    ("6.1 报销时限与流程", [
        "费用发生后30天内需在OA系统提交报销申请，流程为：直属上级审批、财务审核、分管总监审批三级。报销款在审批通过后5个工作日内到账。"]),
    ("6.2 差旅标准", [
        "出差住宿标准：一线城市（北京、上海、广州、深圳）每晚400元，其他城市每晚300元。出差餐补为每天100元，市内交通费实报实销，需注明事由。"]),
    ("6.3 发票要求", [
        "所有报销必须提供合规发票，发票抬头为星辰科技有限公司，税号可在OA首页查询。打车费需逐笔注明起止地点与事由。"]),
    ("第七章 信息安全与保密", []),
    ("7.1 账号与设备安全", [
        "办公电脑必须设置锁屏密码，离开工位需锁屏，系统自动锁屏时间为5分钟。密码须包含大小写字母、数字与特殊字符，每90天更换一次。"]),
    ("7.2 保密义务", [
        "员工不得将公司代码、文档、客户数据上传至个人网盘或社交平台。保密义务在离职后持续2年。核心岗位员工需另行签署竞业限制协议。"]),
    ("第八章 离职流程", [
        "转正员工辞职需提前30天提交书面申请，试用期员工提前3天。离职前需完成工作交接清单，归还全部公司设备与门禁卡。IT部门将在最后工作日回收账号权限。"]),
]


def make_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    font = None
    for name, path, kw in [
        ("SimHei", r"C:\Windows\Fonts\simhei.ttf", {}),
        ("MSYaHei", r"C:\Windows\Fonts\msyh.ttc", {"subfontIndex": 0}),
        ("SimSun", r"C:\Windows\Fonts\simsun.ttc", {"subfontIndex": 0}),
    ]:
        try:
            pdfmetrics.registerFont(TTFont(name, path, **kw))
            font = name
            break
        except Exception:  ***REMOVED*** noqa: BLE001
            continue
    if font is None:
        raise RuntimeError("未找到可用的中文字体（simhei/msyh/simsun）")

    h1 = ParagraphStyle("h1", fontName=font, fontSize=16, leading=24, spaceBefore=14, spaceAfter=6)
    h2 = ParagraphStyle("h2", fontName=font, fontSize=13, leading=20, spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("body", fontName=font, fontSize=10.5, leading=19, firstLineIndent=21)

    story = [Paragraph("星辰科技员工手册", ParagraphStyle("title", fontName=font, fontSize=20, leading=28, spaceAfter=12)),
             Paragraph("（2025年修订版 · 仅供内部使用）", body), Spacer(1, 6 * mm)]
    for title, paras in HANDBOOK:
        style = h1 if title.startswith("第") else h2
        story.append(Paragraph(title, style))
        for p in paras:
            story.append(Paragraph(p, body))

    doc = SimpleDocTemplate(str(OUT / "星辰科技员工手册.pdf"), pagesize=A4,
                            topMargin=20 * mm, bottomMargin=20 * mm)
    doc.build(story)
    print(f"  ✅ 星辰科技员工手册.pdf")


***REMOVED*** ==================== 2. 产品说明书 Word ====================

def make_docx():
    import docx

    d = docx.Document()
    d.add_heading("云滴智能咖啡机 C3 用户手册", 0)
    d.add_paragraph("版本 V2.1 · 云滴电器有限公司")

    d.add_heading("产品概述", 1)
    d.add_paragraph("云滴C3是一款全自动研磨一体式咖啡机，支持美式、意式浓缩、拿铁等多种杯型，"
                    "采用意大利进口高压泵，一键出杯，适合家庭与小型办公室使用。")

    d.add_heading("规格参数", 1)
    table = d.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "项目"
    hdr[1].text = "参数"
    specs = [
        ("型号", "C3"),
        ("额定电压", "220V~ / 50Hz"),
        ("额定功率", "1350W"),
        ("泵压", "20bar（意大利进口高压泵）"),
        ("水箱容量", "1.5L"),
        ("豆仓容量", "250g"),
        ("研磨档位", "共9档（1档最细，9档最粗）"),
        ("整机尺寸", "280×380×420mm"),
        ("净重", "7.8kg"),
        ("待机功耗", "小于1W"),
        ("工作噪音", "小于58dB"),
    ]
    for k, v in specs:
        row = table.add_row().cells
        row[0].text = k
        row[1].text = v

    d.add_heading("快速上手", 1)
    d.add_heading("首次使用", 2)
    d.add_paragraph("首次使用前，请用清水清洗水箱并注水至MAX刻度线，空烧两杯并弃掉，以去除生产残留。")
    d.add_heading("日常制作", 2)
    d.add_paragraph("放入咖啡豆，选择杯量（单杯/双杯），一键启动。预热时间约45秒，出杯完成后自动进入待机状态。")

    d.add_heading("清洁与保养", 1)
    d.add_paragraph("每日倾倒并清洗渣盒；每周清洗水箱与接水盘。建议每2个月进行一次除垢："
                    "长按除垢键3秒进入除垢模式，使用食品级柠檬酸按说明书比例溶解后执行。"
                    "冲泡器滤网每月用温水刷洗一次，请勿使用洗洁精。")

    d.add_heading("故障排除", 1)
    faults = [
        "指示灯红色闪烁：水箱缺水，请注水至MAX线。",
        "不出咖啡：检查水箱是否缺水、豆仓是否已空；若研磨过细请调粗研磨档位。",
        "咖啡口感过淡：将研磨档位调小（更细），或增加粉量设置。",
        "显示屏出现E3：除垢提醒，请尽快执行除垢程序。",
        "机器底部漏水：检查水箱是否安装到位，接水盘是否已溢出。",
    ]
    for f in faults:
        d.add_paragraph(f, style="List Bullet")

    d.add_heading("保修与售后", 1)
    d.add_paragraph("整机自购买之日起保修2年，研磨电机保修3年。人为损坏、未按说明书使用导致的故障不在保修范围内。"
                    "售后热线400-800-1234，报修时需提供购买凭证。")

    d.save(str(OUT / "云滴咖啡机C3产品说明书.docx"))
    print("  ✅ 云滴咖啡机C3产品说明书.docx")


***REMOVED*** ==================== 3/4. TXT 文档 ====================

REMOTE_WORK = """星辰科技远程办公管理制度

一、总则
本制度适用于公司全体正式员工。员工每周申请远程办公不得超过2天，同一部门远程人数不超过部门人数的三分之一。远程办公需提前1个工作日在OA系统提交申请，经直属上级审批后生效。

二、设备与信息安全
远程办公期间使用公司配发的笔记本电脑办公，必须通过公司VPN访问内部系统。禁止在公共Wi-Fi环境下访问内网核心业务系统。视频会议需开启虚拟背景，防止家庭环境信息泄露。办公设备遗失或被盗，须在2小时内上报IT部门与直属上级。

三、工作规范
远程办公期间工作时间与坐班一致，为9:00-18:00。每日9:30的部门站会必须以视频形式参加。工作时间保持企业微信在线，消息响应时间不超过30分钟。远程期间的工作任务量与考核标准不变。

四、考勤与补贴
远程办公当天无需打卡，按正常出勤记录。公司按每月100元标准发放远程办公补贴，用于补偿电费与网络费用，随工资发放。

五、违规处理
未经申请擅自远程办公的，当天按事假处理；一个月内违规累计3次的，取消当月剩余远程办公资格，并通报部门负责人。
"""

IT_FAQ = """星辰科技IT服务台常见问题（FAQ）

Q1：忘记OA系统密码怎么办？
答：可在OA登录页点击"忘记密码"，通过绑定手机号自助重置；也可携带工牌到3楼IT服务台，由管理员人工重置。

Q2：如何申请VPN权限？
答：在OA系统提交"VPN开通申请"，注明使用事由，经部门负责人审批后，IT部门将在1个工作日内开通。

Q3：办公Wi-Fi密码在哪里查看？
答：登录OA首页，在"IT服务公告"栏目查看最新一期Wi-Fi密码，密码每季度更换一次。

Q4：电脑出现硬件故障如何报修？
答：拨打IT服务台电话（内线8888）或在工单系统提交报修申请，描述故障现象并附上截图，IT工程师将在半个工作日内响应。

Q5：门禁卡丢失如何补办？
答：携带工牌到一楼行政前台办理补办，工本费50元，从当月工资中代扣，旧卡即时挂失失效。

Q6：如何申请安装新软件？
答：在OA系统提交"软件安装申请"，注明软件名称与用途；涉及付费软件的需部门负责人审批预算，由IT部门统一安装。
"""


def make_txt():
    (OUT / "远程办公管理制度.txt").write_text(REMOTE_WORK, encoding="utf-8")
    (OUT / "IT服务常见问题FAQ.txt").write_text(IT_FAQ, encoding="utf-8")
    print("  ✅ 远程办公管理制度.txt")
    print("  ✅ IT服务常见问题FAQ.txt")


if __name__ == "__main__":
    print(f"生成示例文档到 {OUT}")
    make_pdf()
    make_docx()
    make_txt()
    print("完成。")
